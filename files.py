"""Incoming chat attachments and files the team writes back to the user."""

from __future__ import annotations

import csv
import io
import logging
import re
import shutil
import sys
import tempfile
import threading
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

# Chainlit adds this directory to sys.path while loading the app, then pops it.
# Keep a copy so lazy imports (and CrewAI re-imports) still find local modules.
_APP_DIR = str(Path(__file__).resolve().parent)
if sys.path[-1:] != [_APP_DIR]:
    sys.path.append(_APP_DIR)

from research import research_tools

logger = logging.getLogger("aiEmployees.files")

MAX_UPLOAD_BYTES = 12 * 1024 * 1024
MAX_TEXT_CHARS = 80_000
MAX_SHEET_ROWS = 80
MAX_SHEET_COLS = 20
MAX_WRITES = 6
MAX_WRITE_BYTES = 2 * 1024 * 1024
SAFE_NAME = re.compile(r"[^A-Za-z0-9._-]+")

WRITE_EXTENSIONS = {
    ".csv",
    ".tsv",
    ".txt",
    ".md",
    ".json",
    ".html",
    ".xml",
    ".py",
    ".sql",
    ".svg",
    ".xlsx",
}

IMAGE_TYPES = {
    "image/jpeg",
    "image/png",
    "image/gif",
    "image/webp",
    "image/jpg",
}

VISION_MARKERS = (
    "gpt-4o",
    "gpt-4.1",
    "gpt-5",
    "gpt-6",
    "claude-3",
    "claude-4",
    "claude-sonnet",
    "claude-opus",
    "claude-fable",
    "claude-mythos",
    "claude-haiku",
    "gemini",
)


def _safe_name(name: str) -> str:
    base = Path(name or "file").name
    cleaned = SAFE_NAME.sub("_", base).strip("._") or "file"
    return cleaned[:80]


def _read_bytes(element) -> bytes | None:
    path = getattr(element, "path", None)
    if path and Path(path).is_file():
        data = Path(path).read_bytes()
        if len(data) > MAX_UPLOAD_BYTES:
            return None
        return data
    content = getattr(element, "content", None)
    if isinstance(content, bytes) and content and len(content) <= MAX_UPLOAD_BYTES:
        return content
    if isinstance(content, str) and content:
        encoded = content.encode("utf-8", errors="replace")
        if len(encoded) <= MAX_UPLOAD_BYTES:
            return encoded
    return None


def _preview_csv(raw: bytes) -> str:
    text = raw.decode("utf-8-sig", errors="replace")
    sample = text[: MAX_TEXT_CHARS * 2]
    try:
        dialect = csv.Sniffer().sniff(sample[:4096], delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(sample), dialect)
    rows = []
    for i, row in enumerate(reader):
        if i >= MAX_SHEET_ROWS:
            rows.append(["…"])
            break
        rows.append([str(cell)[:120] for cell in row[:MAX_SHEET_COLS]])
    if not rows or not rows[0]:
        return "(empty spreadsheet)"
    widths = [max(len(r[c]) if c < len(r) else 0 for r in rows) for c in range(len(rows[0]))]
    lines = []
    for row in rows:
        padded = [(row[c] if c < len(row) else "").ljust(widths[c]) for c in range(len(widths))]
        lines.append(" | ".join(padded))
    return "\n".join(lines)


def _preview_xlsx(raw: bytes) -> str:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    chunks = []
    for sheet in book.worksheets[:4]:
        lines = [f"# {sheet.title}"]
        for i, row in enumerate(sheet.iter_rows(max_row=MAX_SHEET_ROWS, max_col=MAX_SHEET_COLS, values_only=True)):
            values = ["" if cell is None else str(cell)[:120] for cell in row]
            lines.append(" | ".join(values))
            if i >= MAX_SHEET_ROWS - 1:
                lines.append("…")
                break
        chunks.append("\n".join(lines))
    book.close()
    return "\n\n".join(chunks) or "(empty workbook)"


def _preview_text(raw: bytes) -> str:
    text = raw.decode("utf-8-sig", errors="replace")
    if len(text) > MAX_TEXT_CHARS:
        return text[:MAX_TEXT_CHARS] + "\n… [truncated]"
    return text or "(empty file)"


@dataclass
class IncomingFiles:
    prompt_block: str
    summary: str
    input_files: dict[str, str] = field(default_factory=dict)

    def vision_files(self, model: str) -> dict[str, str]:
        model = (model or "").lower()
        if not self.input_files:
            return {}
        if any(marker in model for marker in VISION_MARKERS):
            return dict(self.input_files)
        return {}


def ingest(elements: list[Any] | None, stash: Path) -> IncomingFiles:
    if not elements:
        return IncomingFiles(prompt_block="", summary="")

    stash.mkdir(parents=True, exist_ok=True)
    blocks: list[str] = []
    names: list[str] = []
    input_files: dict[str, str] = {}

    for index, element in enumerate(elements, start=1):
        name = _safe_name(getattr(element, "name", "") or f"attachment-{index}")
        mime = (getattr(element, "mime", None) or "").lower()
        raw = _read_bytes(element)
        if raw is None:
            blocks.append(f"- {name}: could not be read (missing, empty, or larger than 12 MB).")
            names.append(name)
            continue

        dest = stash / f"{index}_{name}"
        dest.write_bytes(raw)
        names.append(name)
        kind = mime or "application/octet-stream"

        if kind.startswith("image/") or kind in IMAGE_TYPES or dest.suffix.lower() in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
            input_files[name] = str(dest)
            blocks.append(f"- Image `{name}` ({kind or 'image'}, {len(raw)} bytes). Look at the attached file.")
            continue

        suffix = dest.suffix.lower()
        try:
            if suffix in {".xlsx", ".xlsm"} or "spreadsheet" in kind:
                preview = _preview_xlsx(raw)
                blocks.append(f"- Spreadsheet `{name}`:\n```\n{preview}\n```")
            elif suffix in {".csv", ".tsv"} or kind in {"text/csv", "text/tab-separated-values"}:
                preview = _preview_csv(raw)
                blocks.append(f"- Spreadsheet `{name}`:\n```\n{preview}\n```")
            elif kind.startswith("text/") or suffix in {".txt", ".md", ".json", ".py", ".sql", ".html", ".xml", ".svg"}:
                preview = _preview_text(raw)
                blocks.append(f"- File `{name}`:\n```\n{preview}\n```")
            else:
                blocks.append(
                    f"- File `{name}` ({kind}, {len(raw)} bytes) is attached. "
                    "Use the filename when referring to it; binary contents are not inlined."
                )
        except Exception as exc:
            logger.warning("Could not preview %s (%s).", name, exc)
            blocks.append(f"- File `{name}` was saved but could not be previewed: {type(exc).__name__}: {exc}")

    prompt_block = "User attached these files:\n" + "\n".join(blocks)
    summary = "Attached: " + ", ".join(names)
    return IncomingFiles(prompt_block=prompt_block, summary=summary, input_files=input_files)


class FileWorkspace:
    """Per-turn folder the employees can write into, then we attach those files."""

    def __init__(self) -> None:
        self.root = Path(tempfile.mkdtemp(prefix="ai-employees-"))
        self.uploads = self.root / "uploads"
        self.outputs = self.root / "out"
        self.uploads.mkdir()
        self.outputs.mkdir()
        self.created: list[Path] = []
        self._lock = threading.Lock()

    def write(self, filename: str, data: bytes) -> Path:
        if len(self.created) >= MAX_WRITES:
            raise ValueError(f"At most {MAX_WRITES} files can be created in one turn.")
        if len(data) > MAX_WRITE_BYTES:
            raise ValueError(f"Each file must be under {MAX_WRITE_BYTES // (1024 * 1024)} MB.")
        name = _safe_name(filename)
        suffix = Path(name).suffix.lower()
        if suffix not in WRITE_EXTENSIONS:
            raise ValueError(
                "Allowed extensions: " + ", ".join(sorted(WRITE_EXTENSIONS))
            )
        dest = self.outputs / name
        with self._lock:
            dest.write_bytes(data)
            if dest not in self.created:
                self.created.append(dest)
        return dest

    def close(self) -> None:
        shutil.rmtree(self.root, ignore_errors=True)


def employee_tools(workspace: FileWorkspace, on_progress=None):
    from crewai.tools import tool

    def emit(message: str) -> None:
        if not on_progress:
            return
        try:
            on_progress(message)
        except Exception:
            logger.debug("progress hook failed", exc_info=True)

    @tool("write_text_file")
    def write_text_file(filename: str, content: str) -> str:
        """Create a downloadable text file for the user (csv, txt, md, json, html, py, sql, svg).
        filename: name with extension. content: full file body."""
        emit(f"is writing `{filename}`…")
        path = workspace.write(filename, content.encode("utf-8"))
        return f"Saved `{path.name}` ({path.stat().st_size} bytes). It will be attached to your chat message."

    @tool("write_spreadsheet")
    def write_spreadsheet(filename: str, csv_content: str) -> str:
        """Create an Excel workbook for the user from CSV text.
        filename: should end in .xlsx. csv_content: comma-separated rows, first row headers."""
        name = filename if filename.lower().endswith(".xlsx") else f"{filename}.xlsx"
        emit(f"is writing spreadsheet `{name}`…")
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.title = "Sheet1"
        reader = csv.reader(io.StringIO(csv_content))
        rows = 0
        for row in reader:
            sheet.append(row)
            rows += 1
            if rows > 5000:
                break
        buffer = io.BytesIO()
        book.save(buffer)
        path = workspace.write(name, buffer.getvalue())
        return f"Saved `{path.name}` with {rows} rows. It will be attached to your chat message."

    return [write_text_file, write_spreadsheet, *research_tools(on_progress=on_progress)]


def chainlit_elements(paths: list[Path]) -> list:
    import chainlit as cl

    elements = []
    for path in paths:
        if not path.is_file():
            continue
        data = path.read_bytes()
        suffix = path.suffix.lower()
        if suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
            elements.append(
                cl.Image(name=path.name, content=data, display="inline", size="medium")
            )
        else:
            elements.append(cl.File(name=path.name, content=data, display="inline"))
    return elements
